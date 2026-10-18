"""
ThreatFusion CTI - Word DOCX Intelligence Briefing Generator
"""

import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from core.report_generator.report_models import ReportData


class DocxReportGenerator:
    @classmethod
    def generate_docx(cls, report: ReportData) -> bytes:
        doc = Document()

        # Page setup
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        # 1. Header TLP Strip
        p_top = doc.add_paragraph()
        run_tlp = p_top.add_run(f"CLASSIFICATION: {report.classification} // TLP:{report.tlp}")
        run_tlp.font.bold = True
        run_tlp.font.size = Pt(9)
        run_tlp.font.color.rgb = RGBColor(210, 153, 34) if report.tlp == "AMBER" else RGBColor(248, 81, 73)
        p_top.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        # 2. Main Title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(report.title)
        title_run.font.size = Pt(20)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(31, 58, 96)

        # 3. Metadata Table
        meta_table = doc.add_table(rows=2, cols=4)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_data = [
            ("REPORT TYPE", report.report_type),
            ("DATE (UTC)", report.date_generated),
            ("AUTHOR", report.author),
            ("TLP PROTOCOL", f"TLP:{report.tlp}")
        ]

        hdr_cells = meta_table.rows[0].cells
        val_cells = meta_table.rows[1].cells
        for idx, (label, val) in enumerate(meta_data):
            hdr_cells[idx].text = label
            val_cells[idx].text = val
            cls._set_cell_background(hdr_cells[idx], "21262D")
            cls._set_cell_background(val_cells[idx], "F6F8FA")
            hdr_cells[idx].paragraphs[0].runs[0].font.size = Pt(8)
            hdr_cells[idx].paragraphs[0].runs[0].font.bold = True
            val_cells[idx].paragraphs[0].runs[0].font.size = Pt(9)

        doc.add_paragraph()  # Spacing

        # 4. Executive Summary
        h1 = doc.add_heading("1. Executive Summary", level=1)
        h1.runs[0].font.color.rgb = RGBColor(31, 58, 96)
        p_sum = doc.add_paragraph()
        p_sum.add_run(report.executive_summary).font.size = Pt(10)

        # Strategic Impact
        if report.strategic_impact:
            h2 = doc.add_heading("2. Strategic Threat Impact", level=1)
            h2.runs[0].font.color.rgb = RGBColor(31, 58, 96)
            doc.add_paragraph(report.strategic_impact).runs[0].font.size = Pt(10)

        # 5. Timeline
        if report.timeline:
            doc.add_heading("3. Chronological Attack Timeline", level=1)
            t_table = doc.add_table(rows=1, cols=3)
            t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            th = t_table.rows[0].cells
            th[0].text = "Date / Phase"
            th[1].text = "Event"
            th[2].text = "Forensic Details"
            for c in th:
                cls._set_cell_background(c, "21262D")
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                c.paragraphs[0].runs[0].font.size = Pt(9)

            for t_item in report.timeline:
                row = t_table.add_row().cells
                row[0].text = str(t_item.get("date", ""))
                row[1].text = str(t_item.get("event", ""))
                row[2].text = str(t_item.get("description", ""))
                for cell in row:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)

        # 6. MITRE ATT&CK Mappings
        if report.mitre_mappings:
            doc.add_heading("4. MITRE ATT&CK Alignment", level=1)
            m_table = doc.add_table(rows=1, cols=4)
            m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            mh = m_table.rows[0].cells
            mh[0].text = "Technique ID"
            mh[1].text = "Technique Name"
            mh[2].text = "Tactic"
            mh[3].text = "Evidence"
            for c in mh:
                cls._set_cell_background(c, "21262D")
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                c.paragraphs[0].runs[0].font.size = Pt(9)

            for m in report.mitre_mappings:
                row = m_table.add_row().cells
                row[0].text = str(m.get("technique_id", ""))
                row[1].text = str(m.get("technique_name", ""))
                row[2].text = str(m.get("tactic", ""))
                row[3].text = str(m.get("evidence", ""))
                for cell in row:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)

        # 7. IOCs Table
        if report.iocs:
            doc.add_heading("5. Indicators of Compromise (IOCs)", level=1)
            ioc_table = doc.add_table(rows=1, cols=4)
            ioc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            ih = ioc_table.rows[0].cells
            ih[0].text = "Type"
            ih[1].text = "Indicator"
            ih[2].text = "C2"
            ih[3].text = "Score"
            for c in ih:
                cls._set_cell_background(c, "21262D")
                c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
                c.paragraphs[0].runs[0].font.size = Pt(9)

            for ioc in report.iocs:
                row = ioc_table.add_row().cells
                row[0].text = str(ioc.get("ioc_type", ""))
                row[1].text = str(ioc.get("ioc_value", ""))
                row[2].text = "YES" if ioc.get("c2_flag") else "NO"
                row[3].text = str(ioc.get("decay_score", 100))
                for cell in row:
                    cell.paragraphs[0].runs[0].font.size = Pt(9)

        # 8. Threat Hunting Queries
        if report.hunt_queries:
            doc.add_heading("6. Operational Threat Hunting Queries", level=1)
            for h in report.hunt_queries:
                doc.add_paragraph(f"[{h.get('language')}] {h.get('title')}", style="Heading 2")
                p_code = doc.add_paragraph()
                r_code = p_code.add_run(h.get("query_content", ""))
                r_code.font.name = "Courier New"
                r_code.font.size = Pt(8.5)

        # 9. Recommendations
        if report.recommendations:
            doc.add_heading("7. Strategic Mitigations & Recommendations", level=1)
            for rec in report.recommendations:
                p_rec = doc.add_paragraph()
                r_tag = p_rec.add_run(f"[{rec.get('priority', 'MEDIUM')}] {rec.get('title')}\n")
                r_tag.font.bold = True
                r_tag.font.size = Pt(10)
                p_rec.add_run(rec.get("details", "")).font.size = Pt(9.5)
        buffer = io.BytesIO()
        doc.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _set_cell_background(cell, hex_color: str):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)
